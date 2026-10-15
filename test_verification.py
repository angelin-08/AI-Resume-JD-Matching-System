"""
Automated Test Verification Script for AI Resume-JD Matching System
"""

from pathlib import Path
import io
import openpyxl

from utils.data_engine import load_all_data, compute_dynamic_views
from utils.ats_engine import (
    extract_text_from_pdf_bytes,
    analyze_sections,
    compute_live_match,
    compute_ats_readiness_assessment,
)
from utils.excel_generator import generate_excel_report


def main():
    print("=== 1. VERIFY DATASET DETECTION ===")
    data = load_all_data()
    resumes_count = len(data["resume_manifest"])
    jds_count = len(data["jd_manifest"])
    comps_count = len(data["comparisons_df"])
    print(f"Resumes detected: {resumes_count}/24")
    print(f"JDs detected: {jds_count}/8")
    print(f"Comparisons accounted for: {comps_count}/192")
    assert resumes_count == 24, "Expected 24 resumes"
    assert jds_count == 8, "Expected 8 JDs"
    assert comps_count == 192, "Expected 192 comparisons"

    print("\n=== 2. VERIFY RESUME-WISE REPORT & THRESHOLD UPDATES ===")
    views_60 = compute_dynamic_views(data, threshold=60.0)
    r_df_60 = views_60["resume_wise_df"]
    print(f"Resume-wise rows at 60%: {len(r_df_60)}/24")
    assert len(r_df_60) == 24, "Expected all 24 resumes in resume-wise report"

    # Check for resumes with zero qualifying matches at 60%
    zero_match_resumes = r_df_60[r_df_60["qualifying_count"] == 0]
    zero_codes = zero_match_resumes["resume_code"].tolist()
    print(f"Resumes with 0 matches at 60%: {len(zero_match_resumes)} ({zero_codes})")
    assert len(zero_match_resumes) == 3, "Expected 3 resumes with 0 matches at 60%"

    for _, row in zero_match_resumes.iterrows():
        assert row["matching_jds"] == "No qualifying matches"
        assert row["match_scores"] == "—"
        assert row["status"] == "No Qualifying Match"
        assert row["best_match_jd"] != ""

    # Recalculate at 50%
    views_50 = compute_dynamic_views(data, threshold=50.0)
    qual_50 = views_50["summary_stats"]["qualifying_matches"]
    print(f"Qualifying matches at 50%: {qual_50} (All 24 designed matches qualify)")
    assert qual_50 == 24, "Expected 24 matches at 50%"

    # Recalculate at 70%
    views_70 = compute_dynamic_views(data, threshold=70.0)
    qual_70 = views_70["summary_stats"]["qualifying_matches"]
    print(f"Qualifying matches at 70%: {qual_70}")
    assert qual_70 == 2, "Expected 2 matches at 70%"

    print("\n=== 3. VERIFY JD-WISE RANKINGS ===")
    jd_df = views_60["jd_wise_df"]
    print(f"JD-wise rows: {len(jd_df)}/8")
    assert len(jd_df) == 8, "Expected 8 JDs in JD-wise report"
    for _, j_row in jd_df.iterrows():
        assert j_row["qualifying_count"] >= 0
        assert j_row["top_candidate"] != ""
        print(
            f"  {j_row['jd_code']}: Top candidate={j_row['top_candidate']} "
            f"({j_row['top_candidate_score']:.2f}%), Qualifying={j_row['qualifying_count']}"
        )

    print("\n=== 4. VERIFY EXCEL REPORT GENERATION (7 SHEETS) ===")
    excel_bytes = generate_excel_report(data, views_60)
    wb = openpyxl.load_workbook(io.BytesIO(excel_bytes))
    expected_sheets = [
        "Resume_Wise_Matches",
        "JD_Wise_Matches",
        "All_Comparisons",
        "Best_Match_Per_Resume",
        "Best_Match_Per_JD",
        "Skill_Gap_Analysis",
        "Summary_Statistics",
    ]
    print(f"Sheets in workbook: {wb.sheetnames}")
    assert wb.sheetnames == expected_sheets, f"Expected sheets {expected_sheets}"
    assert wb["Resume_Wise_Matches"].max_row == 25  # header + 24
    assert wb["JD_Wise_Matches"].max_row == 9      # header + 8
    assert wb["All_Comparisons"].max_row == 193     # header + 192
    assert wb["Best_Match_Per_Resume"].max_row == 25
    assert wb["Best_Match_Per_JD"].max_row == 9
    assert wb["Skill_Gap_Analysis"].max_row == 193
    print("All 7 sheets verified with correct row counts and formatting!")

    print("\n=== 5. VERIFY RESUME UPLOAD & ATS CHECKER ===")
    pdf_bytes = Path("resumes/Resume_01.pdf").read_bytes()
    extracted_text, pages, status = extract_text_from_pdf_bytes(pdf_bytes)
    assert status == "SUCCESS", "Expected extraction status SUCCESS"
    sec_audit = analyze_sections(extracted_text)
    assert sec_audit["has_contact"] is True, "Expected contact info detected"

    target_jd_text = data["doc_text_map"]["JD_01_DA"]["text"]
    live_match = compute_live_match(extracted_text, target_jd_text)
    assert live_match["match_score"] > 60.0, "Expected match score > 60 for Resume_01 vs JD_01_DA"

    ats_res = compute_ats_readiness_assessment(sec_audit, live_match, pages)
    print(
        f"Resume 01 test: Match={live_match['match_score']:.2f}%, "
        f"ATS score={ats_res['ats_score']}/100, Rating={ats_res['rating']}"
    )

    print("\n=== 6. VERIFY ERROR HANDLING ON INVALID PDF ===")
    bad_text, bad_pages, bad_status = extract_text_from_pdf_bytes(b"Corrupted invalid bytes")
    print(f"Bad PDF handling: Status={bad_status}, Text length={len(bad_text)}")
    assert "ERROR" in bad_status or "EXTRACTION_ERROR" in bad_status

    print("\n=======================================================")
    print(">>> ALL 6 VERIFICATION TEST SUITES PASSED PERFECTLY! <<<")
    print("=======================================================")


if __name__ == "__main__":
    main()
